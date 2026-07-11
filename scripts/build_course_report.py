from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "course-report.md"
ASSETS = ROOT / "docs" / "assets"
OUTPUT = ROOT / "deliverables" / "ClaimScope-course-report.docx"
BLUE = "1F5A7A"
INK = "17212B"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    return ImageFont.truetype(str(next(path for path in candidates if path.exists())), size)


def rounded_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, subtitle: str, fill: str) -> None:
    draw.rounded_rectangle(box, radius=12, fill=fill, outline="#C7D3DA", width=2)
    x1, y1, x2, y2 = box
    draw.text(((x1 + x2) / 2, y1 + 28), title, font=font(25, True), fill="#17212B", anchor="mm")
    draw.multiline_text(((x1 + x2) / 2, y1 + 67), subtitle, font=font(16), fill="#53636D", anchor="mm", align="center", spacing=4)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line((start, end), fill="#577683", width=4)
    x, y = end
    draw.polygon([(x, y), (x - 12, y - 7), (x - 12, y + 7)], fill="#577683")


def build_diagrams() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1800, 470), "#F7FAFB")
    draw = ImageDraw.Draw(canvas)
    titles = ["Research\nDirection", "Core Claim", "Variants &\nBoundaries", "Hidden\nAssumptions", "Evidence\nQueries", "Evidence\nCards", "Idea\nOpportunities"]
    colors = ["#EAF3F5", "#DDEFF1", "#F1EEE4", "#EEF1E7", "#E8EEF4", "#F3EBE7", "#E5F1E9"]
    width, gap, y1, y2 = 210, 38, 145, 325
    for index, (title, color) in enumerate(zip(titles, colors)):
        x1 = 35 + index * (width + gap)
        rounded_box(draw, (x1, y1, x1 + width, y2), title, f"Stage {index + 1}", color)
        if index < len(titles) - 1:
            arrow(draw, (x1 + width + 4, 235), (x1 + width + gap - 7, 235))
    draw.text((50, 52), "ClaimScope: assumption-centric pre-ideation workflow", font=font(34, True), fill="#1F5A7A")
    canvas.save(ASSETS / "report-workflow.png", quality=95)

    canvas = Image.new("RGB", (1600, 920), "#F7FAFB")
    draw = ImageDraw.Draw(canvas)
    draw.text((60, 55), "Core Claim Arena: public artifacts, adversarial review", font=font(34, True), fill="#1F5A7A")
    rounded_box(draw, (650, 130, 950, 260), "Research Direction", "fuzzy user input", "#EAF3F5")
    proposer_boxes = [(90, 350, 440, 510), (625, 350, 975, 510), (1160, 350, 1510, 510)]
    proposer_titles = ["Operationalizer", "Mechanism Analyst", "Skeptical Empiricist"]
    proposer_subtitles = ["variables, baseline, metric", "mechanism, target, boundary", "narrow falsifiable claim"]
    for box, title, subtitle in zip(proposer_boxes, proposer_titles, proposer_subtitles):
        rounded_box(draw, box, title, subtitle, "#DDEFF1")
        draw.line(((800, 260), ((box[0] + box[2]) // 2, box[1])), fill="#577683", width=4)
    critic_boxes = [(260, 610, 660, 760), (940, 610, 1340, 760)]
    for box, title, subtitle in zip(critic_boxes, ["Falsifiability Critic", "Scope Critic"], ["observable adverse test", "bounded population and task"]):
        rounded_box(draw, box, title, subtitle, "#F1EEE4")
    for pbox in proposer_boxes:
        center = ((pbox[0] + pbox[2]) // 2, pbox[3])
        for cbox in critic_boxes:
            draw.line((center, ((cbox[0] + cbox[2]) // 2, cbox[1])), fill="#B2C0C7", width=2)
    rounded_box(draw, (650, 790, 950, 900), "Judge", "selection + open slots", "#E5F1E9")
    for cbox in critic_boxes:
        draw.line((((cbox[0] + cbox[2]) // 2, cbox[3]), (800, 790)), fill="#577683", width=4)
    canvas.save(ASSETS / "report-arena.png", quality=95)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_run_font(run, name: str = "宋体", size: float = 10.5, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_toc(paragraph) -> None:
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "更新目录后显示章节与页码"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 14, INK), ("Heading 3", 12, INK)]:
        style = doc.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("ClaimScope  |  自然语言处理课程大作业")
    set_run_font(run, "微软雅黑", 8.5, color="6A7880")
    add_page_number(section.footer.paragraphs[0])
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()
    properties = doc.core_properties
    properties.title = "ClaimScope：面向科研想法生成前的主张、假设与负面证据发现智能体"
    properties.author = "刘子谦"
    properties.subject = "自然语言处理课程大作业"
    properties.keywords = "科研智能体, 多智能体, 科学主张验证, MCP"


def add_cover(doc: Document, lines: list[str]) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(lines[0].replace("# ", ""))
    set_run_font(run, "微软雅黑", 24, True, BLUE)
    p.paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(lines[1]), "微软雅黑", 16, True, INK)
    for _ in range(5):
        doc.add_paragraph()
    for text in lines[2:4]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(text), "微软雅黑", 11, False, "44515A")
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目录"), "微软雅黑", 18, True, BLUE)
    add_toc(doc.add_paragraph())
    doc.add_page_break()


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [rows[0]] + rows[2:]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = re.sub(r"`([^`]+)`", r"\1", value)
            if row_index == 0:
                set_cell_shading(cell, "DDEFF1")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_run_font(run, "微软雅黑" if row_index == 0 else "宋体", 9, row_index == 0)


def build_report() -> Path:
    build_diagrams()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc, [line for line in lines[:8] if line.strip()][:4])
    index = lines.index("## 摘要")
    in_code = False
    figure_number = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if line.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.7)
            run = p.add_run(raw)
            set_run_font(run, "Consolas", 8.5, color="34444D")
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].replace("|", "").replace(":", "").replace("-", "").strip()) == set():
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        image_match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            figure_number += 1
            image_path = SOURCE.parent / image_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            run = p.add_run()
            run.add_picture(str(image_path), width=Cm(15.8))
            caption = doc.add_paragraph(f"图 {figure_number}　{image_match.group(1)}")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent = Cm(0)
            set_run_font(caption.runs[0], "微软雅黑", 9, color="60717A")
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(line[2:])
        elif line:
            p = doc.add_paragraph()
            for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", line):
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith("`") and part.endswith("`"):
                    run = p.add_run(part[1:-1])
                    set_run_font(run, "Consolas", 9, color="1F5A7A")
                else:
                    p.add_run(part)
        index += 1
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
